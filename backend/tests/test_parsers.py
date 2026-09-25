from pathlib import Path

import docx
import fitz

from app.ingestion.parsers import SUPPORTED_EXTENSIONS, UnsupportedFileType, parse_file


def test_txt_parser_reads_plain_text(tmp_path: Path) -> None:
    file_path = tmp_path / "notes.txt"
    file_path.write_text("Our office is open 10am to 6pm.", encoding="utf-8")

    blocks = parse_file(file_path)

    assert len(blocks) == 1
    assert "10am to 6pm" in blocks[0].text
    assert blocks[0].source_name == "notes.txt"


def test_docx_parser_extracts_paragraphs_headings_and_tables(tmp_path: Path) -> None:
    file_path = tmp_path / "policy.docx"
    document = docx.Document()
    document.add_heading("Leave Policy", level=1)
    document.add_paragraph("Employees get 24 days of paid leave per year.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Role"
    table.rows[0].cells[1].text = "Bonus"
    document.save(str(file_path))

    blocks = parse_file(file_path)
    texts = [b.text for b in blocks]

    assert any("24 days of paid leave" in t for t in texts)
    assert any("Role | Bonus" in t for t in texts)
    paragraph_block = next(b for b in blocks if "24 days" in b.text)
    assert paragraph_block.heading == "Leave Policy"


def test_pdf_parser_extracts_text_per_page(tmp_path: Path) -> None:
    file_path = tmp_path / "handbook.pdf"
    doc = fitz.open()
    for text in ("Page one content about onboarding.", "Page two content about payroll."):
        page = doc.new_page()
        page.insert_text((72, 72), text)
    doc.save(str(file_path))
    doc.close()

    blocks = parse_file(file_path)

    assert len(blocks) == 2
    assert blocks[0].page == 1
    assert "onboarding" in blocks[0].text
    assert blocks[1].page == 2
    assert "payroll" in blocks[1].text


def test_parse_file_rejects_unsupported_extension(tmp_path: Path) -> None:
    file_path = tmp_path / "data.csv"
    file_path.write_text("a,b\n1,2", encoding="utf-8")

    try:
        parse_file(file_path)
        assert False, "expected UnsupportedFileType"
    except UnsupportedFileType:
        pass


def test_supported_extensions_matches_parsers() -> None:
    assert set(SUPPORTED_EXTENSIONS) == {".pdf", ".docx", ".txt"}
