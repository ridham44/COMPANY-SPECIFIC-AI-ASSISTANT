from pathlib import Path

import docx

from app.ingestion.base import ParsedBlock


def parse(file_path: Path) -> list[ParsedBlock]:
    document = docx.Document(str(file_path))
    blocks: list[ParsedBlock] = []
    current_heading: str | None = None

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = (para.style.name or "") if para.style else ""
        if style_name.lower().startswith("heading"):
            current_heading = text
            continue
        blocks.append(ParsedBlock(text=text, source_name=file_path.name, heading=current_heading))

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(
                    ParsedBlock(text=" | ".join(cells), source_name=file_path.name, heading=current_heading)
                )

    return blocks
